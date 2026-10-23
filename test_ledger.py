import io, sqlite3, sys, os
from datetime import date
import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import db, engine
from importer import import_workbook, LEDGER_COLS
from app import create_app


def make_xlsx(rows=None, fira2="FIRA-TEST-0002"):
    wb = Workbook()
    p = wb.active; p.title = "Banking Profile"
    p.append(["Field", "Value"])
    for r in [("Beneficiary Account Name", "Test Co"), ("Core Electronic Account", "000011112222"),
              ("IFSC Central Identifier", "BARB0VJMETA"), ("SWIFT Network Address", "BARBINBBAPE")]:
        p.append(r)
    L = wb.create_sheet("Remittance Ledger"); L.append(LEDGER_COLS)
    rows = rows or [
        ["2026-10-01", "TXN-1001", "Alpha", "Export Revenue", "USD", 1500, 84.2, 84, "=F2*G2", 500, None, None, "P0802", "FIRA-TEST-0001"],
        ["2026-10-05", "TXN-1002", "Beta", "Export Revenue", "EUR", 1200, 91.5, 91, "=F3*G3", 600, None, None, "P0807", fira2],
        ["2026-10-10", "TXN-1003", "Hosting", "Infrastructure", "INR", 0, 1, 1, -15000, 0, None, None, "N/A", "N/A"],
        ["2026-10-15", "TXN-1004", "Subcontractor", "Operations", "INR", 0, 1, 1, -45000, 0, None, None, "N/A", "N/A"],
        ["2026-10-20", "TXN-1005", "Gamma", "Export Revenue", "USD", 2500, 84.15000000000001, 84, "=F6*G6", 750, None, None, "P0802", "FIRA-TEST-0003"],
    ]
    for r in rows:
        L.append(r)
    B = wb.create_sheet("Budget vs Actual"); B.append(["Category", "Budgeted INR", "Actual", "Var", "FX"])
    for c, v in [("Export Revenue", 450000), ("Infrastructure", 200000), ("Operations", 100000),
                 ("Marketing", 50000), ("Contingency", 20000)]:
        B.append([c, v])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


@pytest.fixture
def conn(tmp_path):
    return db.connect(str(tmp_path / "t.db"))


def test_import_matches_hand_calculation(conn):
    rep = import_workbook(conn, make_xlsx())
    assert not rep["errors"] and len(rep["posted"]) == 5
    lines = db.statement(conn)
    assert [engine.fmt_inr(l.running_balance) for l in lines] == [
        "125800.00", "235000.00", "220000.00", "175000.00", "384625.00"]
    assert sum(l.fx_variance for l in lines) == 127500          # 300 + 600 + 375
    bva = engine.budget_vs_actual(lines, db.budgets(conn))
    rows = {r["category"]: r for r in bva["rows"]}
    assert rows["Export Revenue"]["actual_settled"] == 44462500
    assert rows["Export Revenue"]["remaining"] == 537500
    assert rows["Infrastructure"]["utilised"] == 1500000 and rows["Infrastructure"]["remaining"] == 18500000
    assert bva["ties_to_ledger"]
    assert db.reconcile(conn)["status"] == "OK"


def test_float_noise_is_normalised():
    assert engine.to_micro(84.15000000000001) == 84_150_000


def test_reimport_is_idempotent_and_changes_are_rejected(conn):
    import_workbook(conn, make_xlsx())
    again = import_workbook(conn, make_xlsx())
    assert not again["errors"] and again["posted"] == [] and len(again["skipped"]) == 5
    rows = [["2026-10-01", "TXN-1001", "Alpha", "Export Revenue", "USD", 1500, 84.3, 84, "=F2*G2", 500, None, None, "P0802", "FIRA-TEST-0001"]]
    changed = import_workbook(conn, make_xlsx(rows))
    assert any("already posted with different values" in e for e in changed["errors"])


def test_dry_run_writes_nothing(conn):
    rep = import_workbook(conn, make_xlsx(), dry_run=True)
    assert len(rep["posted"]) == 5
    assert conn.execute("SELECT COUNT(*) FROM ledger_entries").fetchone()[0] == 0


def test_compliance_rules(conn):
    bad = [["2026-10-01", "TXN-2001", "x", "Export Revenue", "USD", 100, 84, 84, "=F2*G2", 0, None, None, "P9999", "FIRA-TEST-0009"],
           ["2026-10-02", "TXN-2002", "y", "Export Revenue", "USD", 100, 84, 84, "=F3*G3", 0, None, None, "P0802", "N/A"],
           ["2026-10-03", "TXN-2003", "z", "Nonsense", "INR", 0, 1, 1, -10, 0, None, None, "N/A", "N/A"]]
    rep = import_workbook(conn, make_xlsx(bad))
    assert rep["errors"] and conn.execute("SELECT COUNT(*) FROM ledger_entries").fetchone()[0] == 0


def test_duplicate_fira_rejected(conn):
    rep = import_workbook(conn, make_xlsx(fira2="FIRA-TEST-0001"))
    assert rep["errors"]


def test_append_only_and_reversal(conn):
    import_workbook(conn, make_xlsx())
    with pytest.raises(sqlite3.DatabaseError):
        conn.execute("UPDATE ledger_entries SET fees_paise=0 WHERE txn_id='TXN-1001'")
    with pytest.raises(sqlite3.DatabaseError):
        conn.execute("DELETE FROM ledger_entries WHERE txn_id='TXN-1001'")
    with db.tx(conn):
        db.reverse(conn, "TXN-1005", date(2026, 10, 25), "wrong client")
    lines = db.statement(conn)
    assert lines[-1].running_balance == 17500000
    assert sum(l.fx_variance for l in lines) == 90000
    with pytest.raises(engine.ValidationError):
        with db.tx(conn):
            db.reverse(conn, "TXN-1005", date(2026, 10, 26), "again")


def test_tampering_is_detected(conn):
    import_workbook(conn, make_xlsx())
    conn.execute("DROP TRIGGER ledger_no_update")
    conn.execute("UPDATE ledger_entries SET fees_paise=1 WHERE txn_id='TXN-1001'")
    r = db.reconcile(conn)
    assert r["status"] == "DRIFT" and r["checks"]["first_bad_txn"] == "TXN-1001"


def test_closed_period_blocks_backdating(conn):
    import_workbook(conn, make_xlsx())
    with db.tx(conn):
        db.take_snapshot(conn, date(2026, 10, 31))
    late = [["2026-10-30", "TXN-3001", "late", "Operations", "INR", 0, 1, 1, -100, 0, None, None, "N/A", "N/A"]]
    rep = import_workbook(conn, make_xlsx(late))
    assert any("period closed" in e for e in rep["errors"])
    assert db.reconcile(conn)["status"] == "OK"


def test_api_auth_masking_and_flow(tmp_path, monkeypatch):
    monkeypatch.setenv("VELOCE_API_TOKEN", "secret-token")
    c = create_app(str(tmp_path / "api.db")).test_client()
    assert c.get("/api/ledger").status_code == 401
    h = {"Authorization": "Bearer secret-token"}
    r = c.post("/api/import", data={"file": (make_xlsx(), "ledger.xlsx")}, headers=h)
    assert r.status_code == 200 and r.json["closing_balance_inr"] == "384625.00"
    assert c.get("/api/profile", headers=h).json["Core Electronic Account"] == "********2222"
    r = c.post("/api/entries", headers=h, json={"txn_id": "TXN-1006", "date": "2026-10-28", "description": "Delta",
               "category": "Export Revenue", "currency": "USD", "foreign_amount": "100", "applied_rate": "84.3",
               "expected_rate": "84", "fees_inr": "50", "purpose_code": "P0802", "fira_ref": "FIRA-TEST-0100"})
    assert r.status_code == 201 and r.json["fx_variance_inr"] == "30.00" and r.json["running_balance_inr"] == "393005.00"
    assert c.get("/api/reconcile", headers=h).json["status"] == "OK"
